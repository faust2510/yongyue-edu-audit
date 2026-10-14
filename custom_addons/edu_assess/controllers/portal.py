# -*- coding: utf-8 -*-
"""评估引擎三侧入口 + xlsx 往返（设计稿 §6 xlsx / §8 三侧入口与权限）。

四段：
  1. 门户矩阵（学生填评教 / 教师整班打分）——「行 = 名单快照、列 = 评分项」服务端渲染
  2. 我的（学生看自己）
  3. 教师审阅
  4. 教务 xlsx 导出/导入两段式（先例：edu_portal/controllers/roles.py:2013-2158
     的 /grading/export + /grading/import + /grading/import/commit）

安全口径（设计稿 §8.3，安全审计 F1 的教训）：
  · 门户组 ACL = 0（ir.model.access.csv 一行都不给）⇒ 门户读写一律 sudo() +
    控制器内**显式域收敛**；每个写路由 methods=['POST'] + csrf=True；先判身份，
    判空**直接 303 回首页**（不 fallthrough）；
  · 教师侧读名单/打分用**非 sudo** env（openeducat_core.group_op_faculty 有 ACL），
    让 security/security.xml 的行级规则真正生效 —— 控制器万一漏判，规则也不会让
    一位教师读到/写到别班；
  · 匿名任务（task.anonymous）一律**不写 scorer_id**，只写
    h:<HMAC(盐, task_id + user_id)> 指纹；结果页另有 summary_min 样本门槛。

## 算分口径（硬约束，2026-09-28 引擎收口）

双轨合并 / 扣分制转换**只有一个实现入口**：`EduAssessRoster._item_cell(item)`
（以及 `_group_subtotal(group)` / `_breakdown()`）。⇒ 本文件（以及任何跨项 / 跨任务
汇总）**禁止**自己遍历 `edu.assess.score.value` 求和 —— 扣分制下 `value` 存的是
"扣了多少"，自己加会算出 14 分这种反向结果（第 1 期的真事故）。要汇总就读模型派生字段
（`weighted_total / total / average / grade / complete / missing_items /
missing_tracks / reachable_max`）或调上面三个入口。
"""
import base64
import io
import json
import re

from odoo import SUPERUSER_ID, http
from odoo.http import request

from odoo.addons.edu_schedule.models.identity import (
    faculty_of_user, student_of_user)

# 家长侧同样走「继承 edu_portal 的控制器类 + leaf 覆盖」：既有页面 /my/child 由
# EduParentPortal 提供，本模块覆写它**加**一块「孩子的评估」，而不是另开一套家长门户。
# current_parent() 与它同源（op.parent 反查的唯一口径）—— 本模块不自己写第二份。
from odoo.addons.edu_portal.controllers.roles import (EduParentPortal,
                                                      current_parent)

from ..models.assess import anon_fingerprint

TARGET_LABELS = {'student': '学生', 'faculty': '教师'}
SCORER_LABELS = {
    'student': '学生评教师',
    'teacher': '教师评学生',
    'self': '自评',
    'peer': '互评',
}
STATE_LABELS = {'draft': '草稿', 'open': '进行中', 'closed': '已结束'}
KIND_LABELS = {
    'teaching': '评教', 'student': '学生综合', 'admin': '行政考评',
    'intern': '实习评估', 'other': '其他',
}

#: xlsx 导入时忽略的列（导出时写，导回时不读 —— 总分一律服务端算，防篡改）
IGNORED_HEADERS = ('总分', '加权总分', '等级', '平均分')


# ----------------------------------------------------------------------
# 小工具（与本仓既有控制器同口径，不发明第二套）
# ----------------------------------------------------------------------

def _read_table_file(f):
    """读上传的 CSV/XLSX → (header, rows)；单元格统一为去空白字符串。

    形态照抄 edu_portal/controllers/roles.py:19-46（同一容器、同一 openpyxl）。
    """
    if not f or not f.filename:
        return [], []
    raw = f.read()
    name = (f.filename or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        sheet = book.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(['' if c is None else str(c).strip() for c in row])
    else:
        import csv
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('gbk', errors='replace')
        rows = [[(c or '').strip() for c in r]
                for r in csv.reader(io.StringIO(text))]
    rows = [r for r in rows if any(r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _col_index(header, *names):
    for i, cell in enumerate(header):
        if (cell or '').strip() in names:
            return i
    return None


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ''
    return (row[idx] or '').strip()


def _current_student():
    """当前登录用户对应的学生档案（口径 = identity.student_of_user）。"""
    return student_of_user(request.env)


def _current_faculty():
    """当前登录用户对应的教师档案（口径 = identity.faculty_of_user）。

    与 edu_portal roles.current_faculty 同源 —— 两处都只是同一个公共 helper
    的薄壳（edu_schedule/models/identity.py），本模块不再自带一份实现。
    """
    return faculty_of_user(request.env)


def _is_school_admin():
    """教务处/系统管理员 = 最高权限（与 edu_portal roles.privileged_staff 同口径）。"""
    user = request.env.user
    return (user.has_group('openeducat_core.group_op_back_office_admin')
            or user.has_group('base.group_system'))


def _is_teacher():
    return request.env.user.has_group('openeducat_core.group_op_faculty')


def _taught_batch_ids(fac):
    """本人任教班级（op.session.faculty_id 口径，照 roles.py:_teacher_taught_batches）。"""
    if not fac:
        return []
    return request.env['op.session'].sudo().search(
        [('faculty_id', '=', fac.id)]).mapped('batch_id').ids


def _open_grade_task_count(env):
    """本库「进行中、可整班打分」的评估任务数（**全局口径，不分身份**）。

    它只服务**诚实空态**这一件事：区分
      ① 一个进行中的评估任务都没有（=「还没有进行中的评估任务」）；
      ② 有进行中的任务，但名单里没有你任教班级的学生（=「没有你可以打分的任务」）。
    这两种情形在教师侧的可打分计数上完全一样（都是 0），混成同一句话就是在编造。

    「我到底能不能打分」不看这个数，看名单收敛后的行为
    （EduAssessPortal._review_rows）—— 本函数**不参与**任何权限判定。

    口径三条与 /assess 首页既有的 open_tasks 查询逐字一致：
    state=open（草稿没发布 / 已结束都不能打分）· 模板启用（停用模板的任务
    在首页本来就不可见）· 目标是学生（教师→学生才是整班打分，学生→教师那是评教）。
    """
    return env['edu.assess.task'].sudo().search_count([
        ('state', '=', 'open'),
        ('template_id.state', '=', 'active'),
        ('target_kind', '=', 'student')])


#: xlsx 列头里「区名 / 项名」的分隔符（导出与导入必须逐字一致）
LABEL_SEP = ' / '


def _task_items(task, include_inactive=False):
    """本任务的列集合 = 模板的**叶子项**（分组节点不是列），请求时读、按 sequence 排序。

    ⚠ 2026-09-28 修两个口径问题（同一函数）：
      ① 此前用 `task.template_id.item_ids` —— 它**包含分组节点**：三层模板会让门户
         矩阵多出"区/分类"列，而给分组节点打分会被模型 `_check_value` 直接
         ValidationError 拒绝（页面等于给了个打不了分的地方、还占一列）。现在走模型
         自己的 `task._leaf_items()`（"分组节点不是列"的唯一实现）。
      ② 默认仍**只取启用项**：`_leaf_items()` 含历史停用项（汇总要用它：停用项的
         历史分必须继续计入总分），但"填新分"的面（矩阵 / xlsx 导入）不该再出现
         已停用列 —— 口径 = "停用后新打分不再出现该列，历史得分仍保留可查"。
         需要历史列时显式 `include_inactive=True`。

    ⚠ 引擎口径（硬约束，见 models/assess.py 顶部）：跨项/跨任务汇总**禁止**自己遍历
    `edu.assess.score.value`（扣分制下它存的是"扣了多少"）。本文件一律只读派生字段
    （`roster._item_cell/_item_average/weighted_total/total/grade/missing_*/reachable_max`）
    或调 `_item_cell` / `_group_subtotal` / `_breakdown` 三个唯一入口。
    """
    items = task._leaf_items()
    if not include_inactive:
        items = items.filtered(lambda i: i.active)
    return items


def _item_label(item):
    """xlsx 列头标签 = **祖先链（区名）前缀** + 项名，如「厨房 / 窗户有灰尘/污渍」。

    为什么必须带前缀：同一模板的不同区里可以有**同名检查项**（《卫生检测表》5 个区里
    "窗户有灰尘/污渍"×3、"墙面有污渍"×2、"垃圾未倒"×2）。只写项名 ⇒ 导出后再导回时
    按名字对列有歧义（实测：这张表导不回去，同一格的分会被记到另一个区的项上）。
    模型层已把唯一键收窄到「同一上级内唯一」，这里把"同一上级"显式写进列头，
    导出 ↔ 导入就成了一一对应。
    深度 1 的项没有祖先 ⇒ 标签 = 项名（与旧文件逐字一致，向后兼容）。
    """
    names = [item.name or '']
    node = item.parent_id
    seen = {item.id}
    while node and node.id not in seen:
        seen.add(node.id)
        names.insert(0, node.name or '')
        node = node.parent_id
    return LABEL_SEP.join(n for n in names if n)


def _item_chain(item):
    """根 → 叶的节点链（含 item 自己；成环时安全退出，不无限上溯）。"""
    # ⚠ 初始 seen 必须是空集：把 item.id 预置进去会让循环条件当场为假、返回空链
    #   （实测踩过：表头模型整条空 ⇒ 页面只剩余一个 <thead>）。
    nodes, node, seen = [], item, set()
    while node and node.id not in seen:
        seen.add(node.id)
        nodes.insert(0, node)
        node = node.parent_id
    return nodes


def _matrix_header(items):
    """矩阵表头模型：列 = 叶子项（与矩阵的列**逐一对应**），行 = 祖先层级。

    为什么必须有它（2026-09-29 补的门户缺口）：三层模板里**跨区同名项**是常态
    （《卫生检测表》5 个区里「窗户有灰尘/污渍」×3、「墙面有污渍」×2、「垃圾未倒」×2），
    页面矩阵此前只写裸项名 ⇒ 三列长得一模一样，人眼分不出哪一列属于哪个区。
    xlsx 侧早已用「区名 / 项名」列头解决（见 _item_label），页面侧此前没接。

    返回 list[list[dict]]，每个单元格 {'label', 'colspan', 'rowspan',
    'is_item', 'title'}：

      · **叶子项名**一律 is_item=True —— 文本仍是**裸项名**（既有断言按裸名匹配
        表头：tests/smoke/eval_rebuild.py 的 V1-④/V2b-②、assess_task_portal.py 的
        H1.1 都按 <th>项名</th> 匹配，改文本会连带打红无关门禁）；
      · **分组（区/分类）** is_item=False —— 文本 = 该层节点名，由视图渲染成
        <th scope="colgroup" colspan=N>：读屏用户靠它拿到分组语义（跨区同名项正是
        这个功能的理由）。H1.1 的本意是「分组节点不得变成**评分列**」，不是「表头里
        不许出现 th」—— 队长 2026-09-29 裁定按本意把 H1.1 改成只数评分列（见
        tests/smoke/assess_task_portal.py 的注释），分组格因此可以回到语义正确的 th。
      · **扁平模板逐字节等价**：全是深度 1 的项 ⇒ 只有一行、每格 colspan=rowspan=1，
        渲染结果与改动前**逐字节一致**（页面无变化，见新验收脚本的 M-③ 控制组）。
      · 列序**照抄** _task_items（= xlsx 导出的同一列序，不许在这里重排：重排会连带
        改掉 xlsx 列序与既有往返断言）。项序若不是「同区相邻」，同一区会出现多个
        分组格 —— 那是如实呈现，不猜、不合并不同区的列。
    """
    items = list(items)
    if not items:
        return []
    chains = [_item_chain(i) for i in items]
    depth = max(len(c) for c in chains)
    rows = []
    for r in range(depth):
        row = []
        k = 0
        while k < len(items):
            chain, item = chains[k], items[k]
            name_row = len(chain) - 1        # 项名所在行 = 它自己的深度 - 1
            if name_row == r:
                row.append({'label': item.name or '', 'colspan': 1,
                            'rowspan': depth - r, 'is_item': True,
                            'title': item.help or ''})
                k += 1
                continue
            if name_row < r:
                # 项名已经在上面的行里（带 rowspan 占位）⇒ 本行不再占列
                k += 1
                continue
            node = chain[r]
            j = k
            while (j < len(items) and len(chains[j]) > r
                   and chains[j][r].id == node.id):
                j += 1
            row.append({'label': node.name or '', 'colspan': j - k,
                        'rowspan': 1, 'is_item': False, 'title': ''})
            k = j
        rows.append(row)
    return rows


def _item_row_groups(items):
    """行=项 表格的**行分组模型**：_matrix_header 的行轴镜像（同一棵树 _item_chain）。

    为什么有它（2026-09-29 补的门户缺口）：/assess/my 的「评分项 / 得分」表是
    **行 = 叶子项**。上一轮只把**列轴**（矩阵表头）做了区名分组，这张行表仍是裸项名
    ⇒ 《卫生检测表》那类跨区同名项（窗户有灰尘/污渍 ×3）在页面上长得一模一样
    （实测：改前两行的 DOM 逐字相同）。

    口径（与 _matrix_header 逐条对应，**不另立第二套**）：
      · 一棵树：分组信息只来自 _item_chain —— 页面不猜、不从项名里抠字；
      · 一个格：分组行渲染成 <th scope="rowgroup" rowspan=N>。列轴的分组格是
        scope="colgroup" + colspan（见 _matrix_header），行轴的镜像就是
        rowgroup + rowspan —— 读屏用户靠它拿到「这几行同属一组」的语义；
      · 一条文本口径：链文字 = 祖先名以 LABEL_SEP 连接（厨房 / 清洁），与 xlsx
        列头 _item_label 同一棵树、同一分隔符；
      · 项名格**仍是裸项名**（既有断言按裸名匹配表头/表体，改文本会连带打红
        无关门禁）；
      · 只合并**连续**同链的行 —— 项序不是「同区相邻」时同一区会出现多个行组格，
        那是如实呈现，不猜、不合并不相邻的组（与 _matrix_header 同一条纪律）；
      · **扁平模板（全是深度 1 的项）不成组**：chain 全为空串 ⇒ 视图不加分组列，
        渲染与改动前同形（控制组见 tests/smoke/assess_my_item_groups.py 的 C 段）。

    返回等长于 items 的 list[dict]：
      {'item': 条目, 'chain': 链文字（扁平模板 / 无祖先时为空串或 '—'）,
       'rowspan': N（组首行，N = 该组行数）或 0（被上面的 rowspan 覆盖，不渲染格）}
    """
    items = list(items)
    if not items:
        return []
    chains = [_item_chain(i) for i in items]
    if max(len(c) for c in chains) < 2:
        # 扁平模板：一个分组格都不生成（视图据此不加列）
        return [{'item': i, 'chain': '', 'rowspan': 0} for i in items]
    keys = [tuple(n.id for n in c[:-1]) for c in chains]
    out, k = [], 0
    while k < len(items):
        key = keys[k]
        j = k + 1
        if key:                      # 无祖先的行自成一组（连续无链行不合并）
            while j < len(items) and keys[j] == key:
                j += 1
        label = LABEL_SEP.join(n.name or '' for n in chains[k][:-1]
                               if n.name) or '—'
        out.append({'item': items[k], 'chain': label, 'rowspan': j - k})
        for n in range(k + 1, j):
            out.append({'item': items[n], 'chain': label, 'rowspan': 0})
        k = j
    return out


def _roster_label(roster):
    subject = roster.faculty_id or roster.student_id
    return subject.name or ''


def _scorer_ref_for(task, user):
    """打分人指纹：匿名任务 = HMAC 指纹；实名任务 = u:<user_id>。"""
    if task.anonymous:
        return anon_fingerprint(request.env, task.id, user.id)
    return 'u:%d' % user.id


def _score_map(task, scorer_ref, field='value'):
    """{roster_id: {item_id: 值}}（只取该打分人自己的记录）。"""
    Score = request.env['edu.assess.score'].sudo()
    data = {}
    for score in Score.search([('task_id', '=', task.id),
                               ('scorer_ref', '=', scorer_ref)]):
        data.setdefault(score.roster_id.id, {})[score.item_id.id] = (
            score.value if field == 'value' else (score.note or ''))
    return data


def _parse_float(raw):
    try:
        return float(str(raw).strip())
    except (TypeError, ValueError):
        return None


def _score_env(env, anonymous):
    """打分记录写入用的环境。

    ⚠ 匿名任务必须**以系统用户身份写库**：Odoo 会给每条记录自动写
    create_uid / write_uid，而 sudo() 只提权、不改 uid —— 若照直写，打分人的
    user_id 就明文躺在每一行里，"匿名"当场失效。用 with_user(SUPERUSER_ID)
    后这些列统一是系统用户，反查不出是谁（评语文本仍由打分人自愿填写）。
    """
    Score = env['edu.assess.score']
    if anonymous:
        return Score.with_user(SUPERUSER_ID).sudo()
    return Score.sudo()


def _write_scores(task, roster, scorer_ref, uid_or_none, cells, notes=None):
    """按唯一键 upsert 一批格子；返回 (写入数, 跳过数)。

    唯一键 = (task_id, item_id, roster_id, scorer_ref) —— 与数据库
    models.Constraint 同源 ⇒ 重复提交不翻倍（幂等）。
    """
    Score = _score_env(request.env, task.anonymous)
    written = skipped = 0
    for item, value in cells.items():
        if value is None:
            skipped += 1
            continue
        value = max(0.0, min(value, item.max_score or 0.0))
        note = (notes or {}).get(item.id)
        vals = {'value': value}
        if note is not None:
            vals['note'] = note[:500]
        rec = Score.search([
            ('task_id', '=', task.id),
            ('item_id', '=', item.id),
            ('roster_id', '=', roster.id),
            ('scorer_ref', '=', scorer_ref)], limit=1)
        if rec:
            rec.write(vals)
        else:
            vals.update({
                'task_id': task.id,
                'item_id': item.id,
                'roster_id': roster.id,
                'scorer_ref': scorer_ref,
                # 匿名任务一律不写实名打分人（校长可见承诺的实现基础）
                'scorer_id': (False if task.anonymous else uid_or_none),
            })
            Score.create(vals)
        written += 1
    return written, skipped


class EduAssessPortal(http.Controller):

    # ------------------------------------------------------------------
    # 0. 入口总页
    # ------------------------------------------------------------------

    @http.route('/assess', type='http', auth='user', website=True)
    def assess_home(self, **kw):
        """评估总入口：按身份给不同的下一步（学生填评教 / 教师审阅 / 教务看结果 / 家长看孩子）。"""
        env = request.env
        stu = _current_student()
        fac = _current_faculty()
        can_grade = _is_school_admin() or _is_teacher()
        Task = env['edu.assess.task'].sudo()
        open_tasks = Task.search([
            ('state', '=', 'open'),
            ('template_id.state', '=', 'active')], order='date_due, id')
        mine = [t for t in open_tasks
                if t.target_kind == 'faculty' and t.scorer_kind == 'student']
        teach = []
        if fac:
            batches = _taught_batch_ids(fac)
            for task in open_tasks:
                if task.target_kind != 'student':
                    continue
                if Task.browse(task.id).roster_ids.filtered(
                        lambda r: r.batch_id.id in batches):
                    teach.append(task)
        # 「整班打分」**常驻入口**的状态（2026-09-29；来由见 assess_portal.xml 的卡片注释）：
        #   · can_grade 与目标路由 /assess/review 的门禁**同一句判据** —— 卡片上出现
        #     的入口一定点得开（旧判据用 faculty=有 op.faculty 档案，教务组没有这份档案
        #     ⇒ 他们**永远看不到这个入口**，用户就是这么找不到它的）；
        #   · grade_open_count 与审阅队列共用 _review_rows（同一份实现）⇒ 卡片上的数字
        #     与点进去看到的列表是同一批任务，不会互相打脸；
        #   · grade_any_open 只区分两种空态（全库一个进行中的评估都没有 vs 有但不是你的班），
        #     见 _open_grade_task_count —— 不参与权限判定。
        grade_rows = self._review_rows() if can_grade else []
        grade_open = [r for r in grade_rows if r['open']]
        return request.render('edu_assess.edu_assess_home', {
            'title': '评估',
            'page_name': 'assess_home',
            'student': stu,
            'faculty': fac,
            # 家长：入口挂在既有「我的孩子」页（孩子的评估历程在那里），见 §7
            'parent': current_parent(),
            'is_admin': _is_school_admin(),
            'fill_tasks': mine if stu else [],
            'grade_tasks': teach,
            'admin_tasks': open_tasks if _is_school_admin() else [],
            'can_grade': can_grade,
            'grade_open_count': len(grade_open),
            'grade_any_open': _open_grade_task_count(env),
        })

    # ------------------------------------------------------------------
    # 1. 学生：填评教（匿名）+ 我的评估结果
    # ------------------------------------------------------------------

    @http.route('/assess/my', type='http', auth='user', website=True)
    def assess_my(self, **kw):
        """我的评估 —— **被评者**视角（学生 **与教师** 都可用）。

        ⚠ 2026-09-28 修（本仓「割裂感」那条链的另一半）：此前这里只认学生
        （`_current_student()` 判空即 303 /my/home），于是**评教任务里的被评教师**
        在门户上根本看不到自己被评的结果 —— 数据层一度也是同一病症：门户的读规则只认
        `student_id.user_id`，而教师被评的名单行 `batch_id` 为空，按 batch 收敛
        永远得 0 行。模型层现在把这条口径收成唯一实现
        （`EduAssessRoster._subject_domain()` / `_my_subject_rows()`，含
        `faculty_id.user_id` 与 `faculty_id.partner_id.user_ids` 两条链），
        控制器**直接接上它，不自己写第二份 domain** —— 否则又会出现"模型修好了、
        页面还是看不到"。

        两分支的边界（有意为之）：
          · 学生分支**逐字保留**原查询与文案（保护既有学生视角断言：泄漏检查、
            隐藏理由"本次评估结果不向学生公开。"、匿名门槛）；
          · 教师分支只保留"我是被评教师"的行，且只取已结束任务（未结束没有可看的结果）；
          · 两个分支都**不用 published_only 预过滤**：模板关掉"结果可见"时给一行
            解释（而不是静静消失），可见性判据在页面这一层真的被执行。
        """
        env = request.env
        stu = _current_student()
        fac = None if stu else _current_faculty()
        if not stu and not fac:
            # 判空直接回首页，不 fallthrough（§8.3 硬口径）
            return request.redirect('/my/home')
        Task = env['edu.assess.task'].sudo()
        Score = env['edu.assess.score'].sudo()
        fill_tasks = []
        if stu:
            for task in Task.search([('state', '=', 'open'),
                                     ('scorer_kind', '=', 'student'),
                                     ('target_kind', '=', 'faculty')],
                                    order='date_due, id'):
                ref = _scorer_ref_for(task, env.user)
                done = Score.search_count([('task_id', '=', task.id),
                                           ('scorer_ref', '=', ref)])
                need = len(task.roster_ids) * len(_task_items(task))
                fill_tasks.append({
                    'task': task,
                    'done': done,
                    'need': need,
                    'pct': min(int(done * 100.0 / need), 100) if need else 0,
                })
        # 我作为被评者的结果（只有模板允许"结果对被评者可见"时才展示明细）
        Roster = env['edu.assess.roster'].sudo()
        rows = []
        if stu:
            mine = Roster.search([('student_id', '=', stu.id),
                                  ('task_id.state', '=', 'closed')])
            public_off = '本次评估结果不向学生公开。'
        else:
            mine = Roster._my_subject_rows(env.user).filtered(
                lambda r: r.faculty_id and r.task_id.state == 'closed')
            public_off = '本次评估结果不向被评者公开。'
        for roster in mine:
            task = roster.task_id
            items = _task_items(task)
            if not task.template_id.score_visible:
                gs = _item_row_groups(items)
                rows.append({'roster': roster, 'items': items,
                             'item_groups': gs,
                             'has_group_col': any(g['chain'] for g in gs),
                             'hidden': True,
                             'reason': public_off, 'cells': {}})
                continue
            gate = (not task.anonymous) or roster.scorer_count >= task.summary_min
            groups = _item_row_groups(items)
            rows.append({
                'roster': roster,
                'items': items,
                # 行分组模型（_item_row_groups = _matrix_header 的行轴镜像）：
                # 表是「行 = 评分项」，跨区同名项靠分组列才分得出是哪一区的。
                # has_group_col = 本模板真有分组（扁平模板为 False ⇒ 视图不加列，
                # 渲染与改动前同形）——判据在控制器算，模板只做输出。
                'item_groups': groups,
                'has_group_col': any(g['chain'] for g in groups),
                'hidden': not gate,
                'reason': '有效样本不足 %d 份，暂不展示明细。' % task.summary_min,
                'cells': {i.id: roster._item_average(i) for i in items},
            })
        return request.render('edu_assess.edu_assess_my', {
            'title': '我的评估',
            'page_name': 'assess_my',
            'student': stu,
            'faculty': fac,
            'fill_tasks': fill_tasks,
            'rows': rows,
            'scorer_labels': SCORER_LABELS,
        })

    @http.route('/assess/fill/<int:task_id>', type='http', auth='user',
                website=True)
    def assess_fill(self, task_id, saved=None, **kw):
        stu = _current_student()
        if not stu:
            return request.redirect('/my/home')
        task = request.env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists() or task.state != 'open' \
                or task.scorer_kind != 'student' \
                or task.target_kind != 'faculty':
            return request.redirect('/assess/my')
        items = _task_items(task)
        ref = _scorer_ref_for(task, request.env.user)
        return request.render('edu_assess.edu_assess_fill', {
            'title': '填写评估',
            'page_name': 'assess_fill',
            'task': task,
            'items': items,
            # 区名分组表头（三层模板跨区同名项靠它区分；扁平模板 = 一行，渲染不变）
            'header_rows': _matrix_header(items),
            'rosters': task.roster_ids,
            'pre_value': _score_map(task, ref, 'value'),
            'pre_note': _score_map(task, ref, 'note'),
            'saved': bool(saved),
            'anonymous': task.anonymous,
            'summary_min': task.summary_min,
        })

    @http.route('/assess/fill/<int:task_id>/save', type='http',
                methods=['POST'], auth='user', website=True, csrf=True)
    def assess_fill_save(self, task_id, **post):
        stu = _current_student()
        if not stu:
            return request.redirect('/my/home')
        env = request.env
        task = env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists() or task.state != 'open' \
                or task.scorer_kind != 'student' \
                or task.target_kind != 'faculty':
            return request.redirect('/assess/my')
        items = {i.id: i for i in _task_items(task)}
        rosters = {r.id: r for r in task.roster_ids}
        ref = _scorer_ref_for(task, env.user)
        pattern = re.compile(r'^v_(\d+)_(\d+)$')
        grouped = {}
        for key, raw in post.items():
            matched = pattern.match(key or '')
            if not matched:
                continue
            rid, iid = int(matched.group(1)), int(matched.group(2))
            if rid not in rosters or iid not in items:
                continue      # 越权/不存在的 id 一律忽略（按不存在处理）
            grouped.setdefault(rid, {})[iid] = _parse_float(raw)
        written = 0
        for rid, cells in grouped.items():
            best = {}
            notes = {}
            for iid, value in cells.items():
                if value is None:
                    continue
                best[items[iid]] = value
                note = post.get('n_%d_%d' % (rid, iid))
                if note is not None:
                    notes[iid] = note
            done, _skipped = _write_scores(
                task, rosters[rid], ref,
                False if task.anonymous else env.uid, best, notes)
            written += done
            if done and rosters[rid].state != 'done':
                # 名单行的 write_uid 同样会记下打分人 ⇒ 匿名任务走系统用户
                target = (rosters[rid].with_user(SUPERUSER_ID).sudo()
                          if task.anonymous else rosters[rid].sudo())
                target.write({'state': 'done'})
        if task.anonymous:
            # 匿名：审计留痕里也不能出现"谁提交的"（/edu/audit 教务可读）
            env['edu.audit.log'].with_user(SUPERUSER_ID).sudo()._log(
                '评估提交', 'edu.assess.task', task.id, task.name or '',
                '匿名评估收到 %d 条打分（按设计不记录提交人）' % written)
        else:
            env['edu.audit.log'].sudo()._log(
                '评估提交', 'edu.assess.task', task.id, task.name or '',
                '%s 提交 %d 条打分' % (stu.name or '', written))
        return request.redirect('/assess/fill/%d?saved=1' % task.id)

    # ------------------------------------------------------------------
    # 2. 教师：整班打分（只限本人任教班）
    # ------------------------------------------------------------------

    def _teacher_roster_domain(self, task):
        """教师可见的名单行域：按本人任教班收敛；教务不受限。

        ⚠ 这里返回的是"控制器口径"，与 security/security.xml 的行级规则
        （rule_edu_assess_roster_faculty_batch）**故意同构**：两层都要写。
        """
        if _is_school_admin():
            return []
        fac = _current_faculty()
        if not fac:
            return [('id', '=', 0)]
        return [('batch_id', 'in', _taught_batch_ids(fac))]

    @http.route('/assess/grade/<int:task_id>', type='http', auth='user',
                website=True)
    def assess_grade(self, task_id, saved=None, error=None, **kw):
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        task = request.env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists() or task.target_kind != 'student':
            return request.redirect('/assess')
        # 非 sudo 读名单：让 ir.rule 真正生效（教师只拿得到自己任教班）
        rosters = request.env['edu.assess.roster'].search(
            [('task_id', '=', task.id)] + self._teacher_roster_domain(task),
            order='batch_id, id')
        items = _task_items(task)
        ref = _scorer_ref_for(task, request.env.user)
        return request.render('edu_assess.edu_assess_grade', {
            'title': '整班打分',
            'page_name': 'assess_grade',
            'task': task,
            'items': items,
            'header_rows': _matrix_header(items),
            'rosters': rosters,
            'pre_value': _score_map(task, ref, 'value'),
            'pre_note': _score_map(task, ref, 'note'),
            'saved': bool(saved),
            'error': error,
            'anonymous': task.anonymous,
        })

    @http.route('/assess/grade/<int:task_id>/save', type='http',
                methods=['POST'], auth='user', website=True, csrf=True)
    def assess_grade_save(self, task_id, **post):
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        env = request.env
        task = env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists() or task.target_kind != 'student':
            return request.redirect('/assess')
        if task.state != 'open':
            return request.redirect(
                '/assess/grade/%d?error=closed' % task.id)
        rosters = {r.id: r for r in env['edu.assess.roster'].search(
            [('task_id', '=', task.id)] + self._teacher_roster_domain(task))}
        items = {i.id: i for i in _task_items(task)}
        ref = _scorer_ref_for(task, env.user)
        pattern = re.compile(r'^v_(\d+)_(\d+)$')
        grouped = {}
        for key, raw in post.items():
            matched = pattern.match(key or '')
            if not matched:
                continue
            rid, iid = int(matched.group(1)), int(matched.group(2))
            if rid not in rosters or iid not in items:
                continue
            grouped.setdefault(rid, {})[iid] = _parse_float(raw)
        written = 0
        for rid, cells in grouped.items():
            best, notes = {}, {}
            for iid, value in cells.items():
                if value is None:
                    continue
                best[items[iid]] = value
                note = post.get('n_%d_%d' % (rid, iid))
                if note is not None:
                    notes[iid] = note
            done, _skipped = _write_scores(
                task, rosters[rid], ref,
                False if task.anonymous else env.uid, best, notes)
            written += done
            if done and rosters[rid].state != 'done':
                rosters[rid].write({'state': 'done'})
        env['edu.audit.log'].sudo()._log(
            '评估打分', 'edu.assess.task', task.id, task.name or '',
            '教师提交 %d 条打分' % written)
        return request.redirect('/assess/grade/%d?saved=1' % task.id)

    # ------------------------------------------------------------------
    # 3. 教师/教务：审阅队列（=「整班打分」的落点）
    # ------------------------------------------------------------------

    def _review_rows(self):
        """可打分任务的审阅行 —— /assess/review 与 /assess 首页卡片**同一份实现**。

        为什么收成一个函数（2026-09-29）：首页卡片上的「有 N 个进行中的评估可以
        整班打分」与点进去看到的列表必须是**同一批任务**，否则卡片与列表会互相打脸
        （两处各写一份 = 多一次漂移机会）。口径逐条照抄改动前的 /assess/review：

          · 任务 state != 'draft'（草稿还没发布，打不了分）；
          · 目标必须是学生（教师 → 学生才是整班打分；学生 → 教师那是评教）；
          · 名单经 _teacher_roster_domain 收敛（教务不过滤、教师只留本人任教班），
            **名单为空的任务不进列表** —— 没有可打分的行，点进去只有一句提示，
            不该被算作「能打分」；
          · open 标记 = 这条任务**现在**能不能真的提交：closed 的任务
            /assess/grade/<id>/save 会拒绝（回 ?error=closed），所以矩阵页不再给它
            「打分」按钮（那正是"点了没反应的按钮"）。
        """
        env = request.env
        rows = []
        for task in env['edu.assess.task'].sudo().search(
                [('state', '!=', 'draft')], order='state, date_due, id'):
            if task.target_kind != 'student':
                continue
            rosters = env['edu.assess.roster'].search(
                [('task_id', '=', task.id)]
                + self._teacher_roster_domain(task))
            if not rosters:
                continue
            done = len(rosters.filtered(lambda r: r.state == 'done'))
            rows.append({
                'task': task,
                'rosters': len(rosters),
                'done': done,
                'open': task.state == 'open',
                'pct': int(done * 100.0 / len(rosters)) if rosters else 0,
            })
        return rows

    @http.route('/assess/review', type='http', auth='user', website=True)
    def assess_review(self, **kw):
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        rows = self._review_rows()
        return request.render('edu_assess.edu_assess_review', {
            'title': '评估审阅',
            'page_name': 'assess_review',
            'rows': rows,
            # 空态是**两句不同的话**（不许把两种情形混成一句，见 _open_grade_task_count）：
            # 有行但没有一条进行中 vs 一条任务都没有 —— 见模板的 yy-assess-review-empty。
            'open_count': len([r for r in rows if r['open']]),
            'any_open': _open_grade_task_count(request.env),
            # 「去新建任务」是教务组/系统组的活（/assess/tasks 的门禁 =
            # roles.privileged_staff）⇒ 对教师不下发这个按钮，免得给个点了被弹回的按钮。
            'is_admin': _is_school_admin(),
            'state_labels': STATE_LABELS,
        })

    # ------------------------------------------------------------------
    # 4. 结果页（含匿名样本门槛）
    # ------------------------------------------------------------------

    def _results_payload(self, task):
        """结果数据 + 门槛判定。

        门槛（校长拍板项 2 的实现基础）：**份数低于门槛时不展示任何结果明细**。
        口径 = (非匿名任务) 或 (答卷份数 ≥ summary_min) 才出明细；
        行内再判一次：该行份数 < 门槛 ⇒ 该行也显示"样本不足"。
        """
        items = _task_items(task)
        rosters = request.env['edu.assess.roster'].search(
            [('task_id', '=', task.id)] + self._teacher_roster_domain(task),
            order='batch_id, id')
        refs = set(task.score_ids.mapped('scorer_ref'))
        task_gate = (not task.anonymous) or len(refs) >= task.summary_min
        rows = []
        for roster in rosters:
            row_gate = task_gate and (
                not task.anonymous or roster.scorer_count >= task.summary_min)
            rows.append({
                'roster': roster,
                'label': _roster_label(roster),
                'batch': roster.batch_id.name or '',
                'gate': row_gate,
                'count': roster.scorer_count,
                'cells': {i.id: roster._item_average(i) for i in items},
                'weighted': roster.weighted_total,
                'average': roster.average,
            })
        return {
            'items': items,
            'rows': rows,
            'task_gate': task_gate,
            'refs': len(refs),
            'hidden_reason': '有效样本 %d 份，低于门槛 %d 份，暂不展示任何明细。'
                             % (len(refs), task.summary_min),
        }

    @http.route('/assess/task/<int:task_id>/results', type='http', auth='user',
                website=True)
    def assess_results(self, task_id, imported=None, skipped=None, **kw):
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        task = request.env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists():
            return request.redirect('/assess')
        payload = self._results_payload(task)
        return request.render('edu_assess.edu_assess_results', {
            'title': '评估结果',
            'page_name': 'assess_results',
            'task': task,
            'items': payload['items'],
            'header_rows': _matrix_header(payload['items']),
            'rows': payload['rows'],
            'task_gate': payload['task_gate'],
            'hidden_reason': payload['hidden_reason'],
            'refs': payload['refs'],
            'anonymous': task.anonymous,
            'summary_min': task.summary_min,
            'can_export': _is_school_admin() or _is_teacher(),
            'imported': imported,
            'skipped': skipped,
        })

    # ------------------------------------------------------------------
    # 5. xlsx 往返（自研两段式）
    #
    # 为什么必须自研导出：教务组**没有** base.group_allow_export（生产闭包实测，
    # 设计稿 §2.5）⇒ Odoo 原生「导出」按钮对教务根本不出现。
    # ------------------------------------------------------------------

    def _xlsx_response(self, workbook, filename):
        buf = io.BytesIO()
        workbook.save(buf)
        data = buf.getvalue()
        return request.make_response(data, headers=[
            ('Content-Type',
             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            ('Content-Disposition',
             "attachment; filename=\"%s\"; filename*=UTF-8''%s"
             % (filename.encode('ascii', 'replace').decode() or 'assess.xlsx',
                base64.b64encode(filename.encode('utf-8')).decode()))])

    @http.route('/assess/template/<int:template_id>/blank.xlsx', type='http',
                auth='user', website=True)
    def assess_blank_xlsx(self, template_id, task_id=None, **kw):
        """导出空表：第 1 行 = 学号/姓名 + 每个启用评分项一列；有 task_id 时带名单快照。

        第 2 张工作表「说明」写填写说明（数据表保持干净，导入只读第一张表）。
        """
        if not _is_school_admin():
            return request.redirect('/my/home')
        env = request.env
        template = env['edu.assess.template'].sudo().browse(template_id)
        if not template.exists():
            return request.redirect('/assess')
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '评估表'
        # 列 = **启用叶子项**（分组节点不是可打分的列；停用项不再出现），
        # 列头 = 区名前缀 + 项名（同名项因此可分）—— 与 _task_items / 导入侧同一份口径
        items = template.item_ids.filtered(
            lambda i: i.item_kind == 'leaf').sorted(lambda i: (i.sequence, i.id))
        target_is_faculty = template.target_kind == 'faculty'
        header = (['姓名'] if target_is_faculty else ['学号', '姓名'])
        header += [_item_label(i) for i in items]
        sheet.append(header)
        task = None
        if task_id:
            task = env['edu.assess.task'].sudo().browse(int(task_id))
            if not task.exists() or task.template_id.id != template.id:
                task = None
        if task:
            for roster in task.roster_ids:
                subject = roster.faculty_id or roster.student_id
                row = ([] if target_is_faculty
                       else [(subject.gr_no or '').strip()])
                row.append(subject.name or '')
                row += [''] * len(items)
                sheet.append(row)
        info = book.create_sheet('说明')
        info.append(['本表由教务系统导出，请勿改动第 1 行表头'
                    '（列头 = 评分项名称；分组的模板写作「区名 / 项名」，同名项靠它区分）。'])
        info.append(['每行一位被评对象；学号为空时按姓名匹配（重名会整行报错）。'])
        info.append(['单元格留空 = 跳过（不会清空系统里已有的分）；超出满分或非数字整行不写入。'])
        info.append(['「总分 / 等级」列由系统计算，导入时不读。'])
        filename = '%s.xlsx' % (template.name or '评估表')
        return self._xlsx_response(book, filename)

    @http.route('/assess/task/<int:task_id>/export.xlsx', type='http',
                auth='user', website=True)
    def assess_export_xlsx(self, task_id, **kw):
        """导出结果：行 = 名单快照，列 = 评分项 + 加权总分 + 平均分。"""
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        task = request.env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists():
            return request.redirect('/assess')
        payload = self._results_payload(task)
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = '评估结果'
        target_is_faculty = task.target_kind == 'faculty'
        header = (['姓名'] if target_is_faculty else ['学号', '姓名'])
        header += [_item_label(i) for i in payload['items']] + ['加权总分', '平均分']
        sheet.append(header)
        for row in payload['rows']:
            subject = row['roster'].faculty_id or row['roster'].student_id
            line = ([] if target_is_faculty
                    else [(subject.gr_no or '').strip()])
            line.append(row['label'])
            for item in payload['items']:
                value = row['cells'].get(item.id)
                line.append('' if value is None else round(value, 2))
            line += [round(row['weighted'], 2), round(row['average'], 2)]
            sheet.append(line)
        info = book.create_sheet('说明')
        info.append(['本表为评估结果导出；「加权总分 / 平均分」由系统计算，导回时不读。'])
        if task.anonymous:
            info.append(['本任务为匿名评估；下表不包含任何打分人身份信息。'])
        filename = '%s-结果.xlsx' % (task.name or '评估')
        return self._xlsx_response(book, filename)

    def _import_context(self, task):
        """导入的列头映射：**标签（区名 / 项名）→ item**，外加"名字唯一时"的裸名。

        三档（顺序 = 优先级）：
          ① 标签精确匹配 —— 导出文件用的就是它，同名项靠区名前缀区分（唯一解）；
          ② 裸项名匹配 —— **只在模板内唯一时**才进映射，用来兼容 2026-09-28 之前
             导出的旧文件（以及《6 张表》里名字本来就唯一的那些）；
          ③ 其余落"未知列头"分支（可选：按列头新建评分项）。
        同名项若只用裸名来匹配，**绝不随便挑一个**：不进映射、进 ③，并在报错里
        点名"有 N 个同名项，请带区名前缀"——静默记到另一个区的项上是最坏的结果。
        """
        by_label, bare = {}, {}
        for item in _task_items(task):
            by_label[_item_label(item)] = item
            bare.setdefault(item.name, []).append(item)
        for name, found in bare.items():
            if len(found) == 1:
                by_label.setdefault(name, found[0])
        return by_label

    def _dup_column_hint(self, task, unknown):
        """未知列头里"其实是同名项"的：点名让用户改用带区名的列头（不猜）。"""
        bare = {}
        for item in _task_items(task):
            bare.setdefault(item.name, []).append(item)
        hints = []
        for head in unknown:
            found = bare.get(head) or []
            if len(found) > 1:
                hints.append('列头「%s」在同一模板下有 %d 个同名项（%s），'
                             '请改用带区名的列头，例如上列之一'
                             % (head, len(found),
                                '、'.join(_item_label(i) for i in found)))
        return hints

    def _resolve_roster(self, task, no, name):
        """按学号优先/姓名兜底找名单行；返回 (roster, error)。"""
        Roster = request.env['edu.assess.roster'].sudo()
        if task.target_kind == 'faculty':
            domain = [('task_id', '=', task.id), ('faculty_id.name', '=', name)]
            matches = Roster.search(domain)
        else:
            matches = Roster.browse()
            if no:
                matches = Roster.search([
                    ('task_id', '=', task.id),
                    ('student_id.gr_no', '=', no)])
            if not matches and name:
                matches = Roster.search([
                    ('task_id', '=', task.id),
                    ('student_id.name', '=', name)])
                if len(matches) > 1:
                    return None, '姓名「%s」在本任务里对应多名被评者，请填学号' % name
        if not matches:
            return None, '本任务名单里没有「%s」' % (no or name)
        if len(matches) > 1:
            return None, '「%s」在本任务里匹配到多行，请填学号' % (no or name)
        return matches, None

    @http.route('/assess/task/<int:task_id>/import', type='http',
                methods=['POST'], auth='user', website=True, csrf=True)
    def assess_import(self, task_id, **post):
        """第一段：解析 xlsx → 预览（valid / invalid 两表 + 旧值→新值）。"""
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        env = request.env
        task = env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists():
            return request.redirect('/assess')
        upfile = request.httprequest.files.get('file')
        header, rows = _read_table_file(upfile)
        items = self._import_context(task)
        by_name = {name: item for name, item in items.items()}
        return self._import_preview(task, header, rows, by_name, upfile,
                                    bool(post.get('create_items')))

    def _import_preview(self, task, header, rows, by_name, upfile, create_items):
        if not header:
            return request.redirect('/assess/task/%d/results' % task.id)
        i_no = _col_index(header, '学号', '学籍号')
        i_name = _col_index(header, '姓名', '学生', '教师')
        unknown = [h for h in header
                   if h and h not in by_name and h not in IGNORED_HEADERS
                   and h not in ('学号', '学籍号', '姓名', '学生', '教师')]
        created_items = []
        if unknown and create_items:
            # 显式勾选「按新列创建评分项」：这是"导题"路径（抄希沃批量导入点评项目）
            Item = request.env['edu.assess.item'].sudo()
            seq = max(task.template_id.item_ids.mapped('sequence') or [0])
            for name in unknown:
                seq += 10
                created_items.append(Item.create({
                    'template_id': task.template_id.id,
                    'name': name,
                    'max_score': 10.0,
                    'weight': 1.0,
                    'sequence': seq,
                }))
            by_name = self._import_context(task)
            unknown = []
        column_map = {}
        for idx, head in enumerate(header):
            if head in by_name:
                column_map[idx] = by_name[head]
        valid, invalid = [], []
        seen_keys = set()
        if unknown:
            reason = '未知列头：%s（默认不新建评分项；如需按新列创建请勾选该选项）' \
                     % '、'.join(unknown)
            hint = self._dup_column_hint(task, unknown)
            if hint:
                reason += '；' + '；'.join(hint)
            return request.render('edu_assess.edu_assess_import', {
                'title': '评估导入',
                'page_name': 'assess_import',
                'task': task,
                'filename': (upfile.filename if upfile else '') or '',
                'preview': {'valid': [], 'invalid': []},
                'file_error': reason,
                'confirm_json': '[]',
                'created_item_names': '、'.join(
                    i.name for i in created_items),
            })
        Score = request.env['edu.assess.score'].sudo()
        uid = request.env.uid
        ref = 'i:%d' % uid
        for n, row in enumerate(rows, start=2):
            no, name = _cell(row, i_no), _cell(row, i_name)
            errors = []
            if task.target_kind == 'faculty':
                if not name:
                    errors.append('缺姓名')
            elif not (no or name):
                errors.append('缺学号与姓名')
            roster, err = (None, None)
            if not errors:
                roster, err = self._resolve_roster(task, no, name)
                if err:
                    errors.append(err)
            cells, notes = {}, {}
            for idx, item in column_map.items():
                raw = _cell(row, idx)
                if raw == '':
                    continue          # 空 = 跳过（不清空库里已有的分）
                value = _parse_float(raw)
                if value is None:
                    errors.append('「%s」不是数字' % item.name)
                    continue
                if value < 0:
                    errors.append('「%s」不能为负' % item.name)
                    continue
                if value > (item.max_score or 0):
                    errors.append('「%s」超过满分 %g'
                                  % (item.name, item.max_score))
                    continue
                cells[item] = value
                # 评语列头：新区名口径（「区名 / 项名_评语」）优先，旧裸名口径兜底
                note_idx = _col_index(header, '%s_评语' % _item_label(item),
                                      '%s_评语' % item.name)
                if note_idx is not None:
                    note = _cell(row, note_idx)
                    if note:
                        notes[item.id] = note
            key = (roster.id if roster else 0, tuple(sorted(
                (i.id, v) for i, v in cells.items())))
            if key in seen_keys:
                errors.append('同一文件里出现了重复的被评对象')
            seen_keys.add(key)
            entry = {
                'row': n,
                'no': no,
                'name': name,
                'roster_id': roster.id if roster else 0,
                'label': _roster_label(roster) if roster else (name or no),
                'cells': [{'item_id': i.id, 'name': i.name, 'value': v}
                          for i, v in sorted(cells.items(),
                                             key=lambda kv: kv[0].sequence)],
                'notes': {str(iid): note for iid, note in notes.items()},
            }
            if errors:
                entry['errors'] = '；'.join(errors)
                invalid.append(entry)
            else:
                if not entry['cells']:
                    entry['empty_row'] = True
                # 预览显示 旧值 → 新值
                for cell in entry['cells']:
                    old = Score.search([
                        ('task_id', '=', task.id),
                        ('roster_id', '=', roster.id),
                        ('item_id', '=', cell['item_id']),
                        ('scorer_ref', '=', ref)], limit=1)
                    cell['old'] = None if not old else old.value
                valid.append(entry)
        payload = [e for e in valid if not e.get('empty_row')]
        return request.render('edu_assess.edu_assess_import', {
            'title': '评估导入',
            'page_name': 'assess_import',
            'task': task,
            'filename': (upfile.filename if upfile else '') or '',
            'preview': {'valid': valid, 'invalid': invalid},
            'file_error': None,
            'confirm_json': json.dumps(payload, ensure_ascii=False),
            'created_item_names': '、'.join(i.name for i in created_items),
        })

    @http.route('/assess/task/<int:task_id>/import/commit', type='http',
                methods=['POST'], auth='user', website=True, csrf=True)
    def assess_import_commit(self, task_id, **post):
        """第二段：确认写入（逐条 savepoint + 唯一键幂等 upsert + 审计留痕）。"""
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        env = request.env
        try:
            items = json.loads(post.get('payload') or '[]')
        except (TypeError, ValueError):
            items = []
        task = env['edu.assess.task'].sudo().browse(task_id)
        if not task.exists() or not isinstance(items, list):
            return request.redirect('/assess')
        Roster = env['edu.assess.roster'].sudo()
        Item = env['edu.assess.item'].sudo()
        Score = env['edu.assess.score'].sudo()
        ref = 'i:%d' % env.uid
        written = skipped = 0
        for entry in items:
            roster = Roster.browse(int((entry or {}).get('roster_id') or 0))
            if not roster.exists() or roster.task_id.id != task.id:
                skipped += 1
                continue
            cells = {}
            notes = {}
            bad = False
            for cell in (entry or {}).get('cells') or []:
                item = Item.browse(int(cell.get('item_id') or 0))
                value = _parse_float(cell.get('value'))
                if not item.exists() or item.template_id.id != task.template_id.id \
                        or value is None or value < 0 \
                        or value > (item.max_score or 0):
                    bad = True
                    break
                cells[item] = value
            if bad or not cells:
                skipped += 1
                continue
            for iid, note in ((entry or {}).get('notes') or {}).items():
                notes[int(iid)] = note
            try:
                with env.cr.savepoint():
                    done, _skip = _write_scores(
                        task, roster, ref, env.uid, cells, notes)
                    written += done
            except Exception:
                skipped += 1
        audit = env['edu.audit.log'].sudo()
        audit._log('评估导入', 'edu.assess.task', task.id, task.name or '',
                   '导入 %d 条 / 跳过 %d 条（文件 %s）' % (
                       written, skipped, post.get('filename') or '—'))
        return request.redirect(
            '/assess/task/%d/results?imported=%d&skipped=%d'
            % (task.id, written, skipped))

    # ------------------------------------------------------------------
    # 5. 学生的评估历程（跨任务聚合；口径与实现说明见文件末段 §6）
    # ------------------------------------------------------------------

    @http.route('/assess/my/history', type='http', auth='user', website=True)
    def assess_my_history(self, **kw):
        """学生本人：历次评估（时间序）。

        **路由上没有任何能指向别人的参数** —— 被评者是 identity 反查出来的
        （student_of_user），所以"看别人的历程"在结构上不可达；即使有人手工
        拼 ?student_id=某人 也不会被读（那只是被 **kw 吃掉的一个普通查询串）。
        """
        stu = _current_student()
        if not stu:
            return request.redirect('/my/home')
        rosters = request.env['edu.assess.roster'].sudo().search(
            _history_domain(stu.id)
            + [('task_id.state', 'in', list(HISTORY_TASK_STATES))])
        return request.render(
            'edu_assess.edu_assess_history',
            dict(title='我的评估历程',
                 **_history_context(True, stu,
                                    _history_entries(rosters, True))))

    @http.route('/assess/student/<int:student_id>/history', type='http',
                auth='user', website=True)
    def assess_student_history(self, student_id, **kw):
        """教师/教务：某个学生的历次评估（从打分矩阵或结果页点名字进来）。

        两层收敛，缺一不可：
          · 控制器层：非教务/教师一律回门户首页；学生/家长连页面都到不了；
          · 数据层：名单行用**非 sudo** env 读 ⇒ security/security.xml 的行级规则
            真正生效（教师只拿得到本人任教班的名单行），控制器万一漏判也不会越界。
        """
        env = request.env
        if not (_is_school_admin() or _is_teacher()):
            return request.redirect('/my/home')
        student = env['op.student'].sudo().browse(student_id)
        if not student.exists():
            return request.redirect('/assess')
        if not self._staff_may_read(student):
            # 不在可见范围：不给数据、也不渲染该生身份（fail-closed）
            return request.render(
                'edu_assess.edu_assess_history',
                dict(title='学生评估历程',
                     **_history_context(False, None, [], out_of_scope=True)))
        rosters = env['edu.assess.roster'].search(
            _history_domain(student_id)
            + [('task_id.state', 'in', list(HISTORY_TASK_STATES))])
        return request.render(
            'edu_assess.edu_assess_history',
            dict(title='学生评估历程',
                 **_history_context(False, student,
                                    _history_entries(rosters, False))))

    @http.route('/assess/faculty/history', type='http', auth='user',
                website=True)
    def assess_my_faculty_history(self, **kw):
        """被评教师本人：历次评教（跨任务、时间序）。

        **路由上没有任何能指向别人的参数** —— 被评者是 identity 反查出来的
        （faculty_of_user）：看别人的教师历程在结构上不可达（与学生侧
        /assess/my/history 同一写法的结构性保证）。

        ⚠ 为什么另开一条、而不接进 /assess/my/history：那条路由是**学生**的，且
        tests/smoke/assess_history.py 的 H5f 明确要求「教师请求学生本人的历程 → 弹回
        门户」（既有契约；本模块只许追加断言、不能改那条）⇒ 教师侧另开一条，页面与
        闸门仍共用同一份实现（_history_entries / _history_context）。
        """
        env = request.env
        fac = _current_faculty()
        if not fac:
            return request.redirect('/my/home')
        # 收敛域走模型唯一口径（faculty_id.user_id ∪ faculty_id.partner_id.user_ids，
        # 见 models/assess.py 的 EduAssessRoster._subject_domain）—— 控制器不自己写
        # 第二份 domain，否则又会「模型修好了、页面还是看不到」。
        mine = env['edu.assess.roster'].sudo()._my_subject_rows(env.user).filtered(
            lambda r: r.faculty_id and r.task_id.state in HISTORY_TASK_STATES)
        return request.render(
            'edu_assess.edu_assess_history',
            dict(title='我的评估历程',
                 **_history_context(
                     True, None,
                     _history_entries(mine, True, subject_word='被评者'),
                     heading='教师评估历程', subject_word='被评者')))

    def _staff_may_read(self, student):
        """教务=全部；教师=该生出现在本人任一任教班的学籍行里（口径与名单快照同源）。

        这里只决定「给不给看这一页」；**真正决定看到几行的是行级规则**（上面用非 sudo
        读名单行）。所以即便本函数判错为 True，越界的行仍然拿不到 —— 两层是纵深，不是替代。
        """
        if _is_school_admin():
            return True
        fac = _current_faculty()
        if not fac:
            return False
        batches = set(_taught_batch_ids(fac))
        if not batches:
            return False
        return bool(student.sudo().course_detail_ids.filtered(
            lambda l: l.batch_id and l.batch_id.id in batches))


# ----------------------------------------------------------------------
# 6. 学生的评估历程（跨任务聚合，**零新模型**）
#
# 为什么有这一段：评估引擎上线后，同一个学生跨任务的历史在**呈现层**是断的 ——
# 分数只活在「某个任务的结果页」里；学生本人在 /assess/my 只看到当前若干张卡，
# 教师也看不到「这个学生历次怎么样」。聚合键用**既有的** roster.subject_ref
# （'s:<id>'，跨任务稳定；口径与字段定义见 models/assess.py 的
# EduAssessRoster.subject_ref —— 按 code-map §6 第 15 条只写符号名，不写行号）
# 与既有的计算字段
# total / weighted_total / average / scorer_count ⇒ 不新增模型、不新增字段、
# 不新增权限面（门户一行 ACL 都没加，读法沿用 sudo() + 显式域收敛）。
#
# 继承既有口径**两条**，不发明第二套：
#   ① 模板的「结果对被评者可见」（score_visible）：只约束被评者本人；
#   ② 匿名任务的样本门槛（summary_min）：两类身份都要过（小样本下匿名承诺才成立）。
#
# 学生视角**只列已结束（closed）的任务** —— 进行中/草稿连"存在"都不列，
# 以免暴露「别人还没打完 / 任务还没截止」的中间态（口径与 /assess/my 的
# "只展示已结束任务" 一致）。教师/教务视角同一条时间口径（历程 = 已经发生的事），
# 进行中的任务他们本来就能在 /assess/review、/assess/grade、结果页按任务查看。
# ----------------------------------------------------------------------

#: 历程页只认**已结束**的任务（口径见上；改这里等于改曝光面，需同步 tests/smoke/assess_history.py）
HISTORY_TASK_STATES = ('closed',)


def _history_domain(student_id):
    """历程查询域 = 稳定键 subject_ref ∩ 真实外键 student_id。

    ⚠ 两道都要钉：subject_ref 是 Char 且**不是**数据库只读列（UI 只读、RPC 可写），
    只按它收敛的话，一旦有人把 student_id=甲 的行写成 subject_ref='s:乙'，
    乙就会在自己的历程页里看到甲的分数。取交集后两列必须**同时**指向同一学生，
    单列写错不再构成越权（判据：tests/smoke/assess_history.py 的 H4/H5）。
    """
    return [('subject_ref', '=', 's:%d' % student_id),
            ('student_id', '=', student_id)]


def _history_order_key(task):
    """时间序的排序键：截止日 → 开始日 → 周期串 → id（缺日期时不靠 id 乱序）。"""
    return (str(task.date_due or task.date_begin or ''),
            task.period or '', task.id)


def _history_entries(rosters, as_subject, subject_word='学生'):
    """名单行（= 一个被评者在各任务里的快照）→ 时间序的历程条目。

    :param as_subject: True = 被评者本人视角（要过 score_visible 那道闸门）；
                       False = 教师/教务视角（不受 score_visible 限制，仍受匿名门槛限制）。
    :param subject_word: 隐藏原因里的被评者称呼（学生 / 被评者 / 孩子）——
                       首参默认值刻意与改动前逐字一致（H7a 断言「不向学生公开」原文）。
    """
    entries = []
    prev_rate = None
    for roster in sorted(rosters, key=lambda r: _history_order_key(r.task_id)):
        task = roster.task_id
        tpl = task.template_id
        has_scores = roster.scorer_count > 0
        sample_ok = (not task.anonymous) or roster.scorer_count >= task.summary_min
        visible = (has_scores and sample_ok
                   and (not as_subject or tpl.score_visible))
        if visible:
            reason = ''
        elif not has_scores:
            reason = '本次评估没有收到任何打分，未生成结果。'
        elif as_subject and not tpl.score_visible:
            reason = '本次评估结果不向%s公开。' % subject_word
        else:
            reason = '有效样本不足 %d 份，暂不展示明细。' % task.summary_min
        full = tpl.total or 0.0
        rate = (roster.average / full * 100.0) if (visible and full) else None
        delta = (rate - prev_rate) if (rate is not None
                                       and prev_rate is not None) else None
        if rate is not None:
            prev_rate = rate
        entries.append({
            'task': task,
            'template': tpl,
            'name': tpl.name or '评估',
            'kind': KIND_LABELS.get(tpl.kind, tpl.kind or ''),
            'period': task.period or '',
            'when': str(task.date_due or task.date_begin or ''),
            'batch': roster.batch_id.name or '',
            # 不可见时**不下发数字**：模板层拿不到就不会误渲染（不是靠"记得别写"）
            'visible': visible,
            'reason': reason,
            'full': full if visible else None,
            'total': roster.total if visible else None,
            'average': roster.average if visible else None,
            'weighted': roster.weighted_total if visible else None,
            'count': roster.scorer_count if visible else None,
            'rate': rate,
            'delta': delta,
        })
    return entries


def _history_trend(entries):
    """首末得分率对比（≥2 条可公开结果才算得出来；不足返回 None，不编造）。"""
    rates = [e['rate'] for e in entries if e['rate'] is not None]
    if len(rates) < 2:
        return None
    return {'n': len(rates), 'first': rates[0], 'last': rates[-1],
            'delta': rates[-1] - rates[0]}


def _history_context(own, student, entries, out_of_scope=False, heading=None,
                     subject_word='学生', back=None, back_label=None):
    """各入口共用一份上下文（键必须齐全：QWeb 少一个键就是一次 500）。

    :param heading: 页标题（教师/家长入口各写各的；学生/教师/教务入口走默认值，
                    默认值 = 改动前的逐字文案，既有断言不受影响）。
    :param subject_word: 被评者称呼（同上，默认值与改动前逐字一致）。
    :param back/back_label: 返回链接与文案（默认 = 改动前的行为）。
    """
    return {
        'page_name': 'assess_history',
        'own': own,
        'student': student,
        'entries': entries,
        'trend': _history_trend(entries),
        'out_of_scope': out_of_scope,
        'heading': heading,
        'subject_word': subject_word,
        'back': back,
        'back_label': back_label,
        'state_labels': STATE_LABELS,
    }


# ----------------------------------------------------------------------
# 7. 家长：自己孩子的评估历程（既有页面 /my/child + 一条历程页）
#
# 家长-孩子绑定的**权威口径**（先核实后动手，取证过程见报告）：
#   · op.parent.student_ids —— 基座 openeducat_parent 的 M2M（库里的表
#     op_parent_op_student_rel：主键 (op_parent_id, op_student_id) + 两条 FK
#     指回 op_parent / op_student）＝ 权威链路；
#   · op.parent.user_id     —— 家长登录账号（门户认家长的唯一字段：
#     edu_portal/models/parent_link_fix.py 的「第 5 条」明写它是规范字段，
#     edu_portal/security/security.xml 的家长行级规则也用 'student_id.parent_ids.user_id'）；
#   · op.student.parent_ids —— 同一张关系表的逆 M2M（不是第二条事实链）；
#   · res.users.child_ids   —— 基座顺带回写的便利列，**不当权威用**（可能为空/过期）。
#   反查 = edu_schedule/models/identity.py 的 profile_of_user('op.parent')（双向并集），
#   与 edu_portal/controllers/roles.py 的 current_parent() 同源 —— 不自己写第二份。
#
# 越权收敛三层（缺一不可）：
#   ① 控制器：孩子 id 必须 ∈ 本人 op.parent.student_ids ⇒ 不在就 303 回 /my/child，
#      连该生身份都不渲染（fail-closed，与既有 /my/child/<id> 同判据）；
#   ② 数据层：edu.assess.* 对 base.group_portal 的 ACL = 0 ⇒ 家长连 RPC 直读都拿不到
#      （判据与 tests/smoke/eval_rebuild.py 的 V8-③ 同款）；
#   ③ 可见性两道闸门与**孩子本人**完全一致（as_subject=True）：模板 score_visible +
#      匿名样本门槛 summary_min；不过时控制器不下发数字（不是「记得别渲染」）。
# ----------------------------------------------------------------------

def _child_history_entries(student_id, subject_word='孩子'):
    """孩子的历程条目（家长视角 = **孩子本人视角**：同一套 domain + 同一套闸门）。"""
    rosters = request.env['edu.assess.roster'].sudo().search(
        _history_domain(student_id)
        + [('task_id.state', 'in', list(HISTORY_TASK_STATES))])
    return _history_entries(rosters, True, subject_word)


def _child_assess_card(student_id, subject_word='孩子'):
    """「我的孩子」卡片上的一小块：全量条目 + 最近 3 条（倒序）+ 两个计数。"""
    entries = _child_history_entries(student_id, subject_word)
    return {'entries': entries,
            'recent': list(reversed(entries))[:3],
            'n': len(entries),
            'visible_n': len([e for e in entries if e['visible']])}


class EduAssessParentPortal(EduParentPortal):
    """家长侧：把孩子收到的评估接进**既有页面** /my/child（不另开一套家长门户）。

    形态 = 继承 edu_portal 的 EduParentPortal 并覆写同一路由（Odoo 路由生成器的
    leaf 覆盖：父类那条不再单独注册；先例 = 本文件末段的 303 兼容层）。覆写只**加**
    数据、不改既有展示：kids 那份摘要仍由父类的 _child_summary() 算 —— 不复制它的口径。
    """

    def _child_summary(self, stu):
        """父类摘要 **+ 一块评估数据**（super() 取父类原值，不重写它的口径）。

        为什么从这一层注入：两个家长页（/my/child 列表、/my/child/<id> 仪表盘）
        都由父类路由渲染，且都调 `self._child_summary(stu)` ⇒ 覆写这一个方法就等于
        **一处**给两页带上 kid['assess']，两条路由的主体逻辑一行都不用碰。
        """
        kid = super()._child_summary(stu)
        kid['assess'] = _child_assess_card(stu.id)
        return kid

    @http.route('/my/child', type='http', auth='user', website=True)
    def my_child(self, **kw):
        """薄壳：**逐字委托**父类（孩子集合 / 摘要口径一行不抄）。

        为什么必须重新声明这条路由：Odoo 的路由生成器按「leaf 覆盖」只实例化最底层的
        子类 —— 子类不重新声明就用不到本类的 `_child_summary`。方法体里 super() 调的
        是父类实现本身，所以这里不存在第二份口径（与上面 303 兼容层同一机制）。
        """
        return super().my_child(**kw)

    @http.route('/my/child/<int:child_id>', type='http', auth='user',
                website=True)
    def my_child_dashboard(self, child_id, **kw):
        """同上：薄壳委托。**仪表盘主体（成绩/摘要/快捷入口）一行不抄** ——
        越权判据也仍是父类那一份（child_id 必须 ∈ 本人 op.parent.student_ids）。"""
        return super().my_child_dashboard(child_id, **kw)

    @http.route('/assess/child/<int:student_id>/history', type='http',
                auth='user', website=True)
    def assess_child_history(self, student_id, **kw):
        """家长：自己孩子的完整评估历程（越权三层见上面那段注释）。"""
        env = request.env
        parent = current_parent()
        stu = env['op.student'].sudo().browse(student_id)
        if not parent or not stu.exists() \
                or student_id not in parent.student_ids.ids:
            return request.redirect('/my/child')
        return request.render('edu_assess.edu_assess_history', dict(
            title='孩子的评估历程',
            **_history_context(False, stu,
                               _child_history_entries(student_id),
                               heading='孩子的评估历程', subject_word='孩子',
                               back='/my/child', back_label='返回我的孩子')))

