
# -*- coding: utf-8 -*-
import sys, openpyxl
from openpyxl.utils import get_column_letter

def dump(path, dest):
    wb = openpyxl.load_workbook(path, data_only=False)
    L = []
    L.append('='*100)
    L.append('FILE: ' + path.split('/')[-1])
    L.append('SHEETS: %d -> %s' % (len(wb.sheetnames), wb.sheetnames))
    L.append('='*100)
    for ws in wb.worksheets:
        L.append('')
        L.append('#'*90)
        L.append('SHEET: %r  dims=%s  max_row=%s max_col=%s  state=%s' % (ws.title, ws.dimensions, ws.max_row, ws.max_column, ws.sheet_state))
        L.append('#'*90)
        merged = sorted(str(m) for m in ws.merged_cells.ranges)
        L.append('MERGED RANGES (%d): %s' % (len(merged), ', '.join(merged) if merged else '(none)'))
        # column widths
        widths = {k: v.width for k,v in ws.column_dimensions.items() if v.width}
        L.append('COL WIDTHS: %s' % widths)
        frz = ws.freeze_panes
        L.append('FREEZE PANES: %s' % frz)
        dv = list(ws.data_validations.dataValidation) if ws.data_validations else []
        if dv:
            L.append('DATA VALIDATIONS:')
            for d in dv:
                L.append('   type=%s formula1=%s sqref=%s' % (d.type, d.formula1, d.sqref))
        # conditional formatting
        try:
            for rng, rules in ws.conditional_formatting._cf_rules.items():
                for r in rules:
                    L.append('COND FMT: %s type=%s op=%s formula=%s' % (rng.sqref, r.type, r.operator, getattr(r,'formula',None)))
        except Exception:
            pass
        L.append('')
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None and not (c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None,'00000000')):
                    continue
                v = c.value
                extra = ''
                if isinstance(v, str) and v.startswith('='):
                    extra = ' [FORMULA]'
                if c.data_type == 'f': extra = ' [FORMULA]'
                nf = c.number_format
                if nf and nf != 'General': extra += ' [fmt:%s]' % nf
                cells.append('%s=%r%s' % (c.coordinate, v, extra))
            if cells:
                L.append('  ' + ' | '.join(cells))
        L.append('')
        L.append('--- END SHEET %r ---' % ws.title)
    open(dest,'w',encoding='utf-8').write('\n'.join(L))
    return '%s -> %s (%d lines)' % (path.split('/')[-1], dest, len(L))

if __name__ == '__main__':
    print(dump(sys.argv[1], sys.argv[2]))
